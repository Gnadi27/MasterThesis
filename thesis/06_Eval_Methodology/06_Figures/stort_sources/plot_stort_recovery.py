#!/usr/bin/env python3
"""Plot STORT wall/recovery-temperature ratios using workbook Mach and T_inf."""
import argparse
from pathlib import Path
import csv
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import load_workbook

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--workbook', type=Path, default=Path('/home/mgnadlinger/ZPG_FP_overview/STORT/STORT_experiment_data.xlsx'))
parser.add_argument('--output-dir', type=Path, default=Path(__file__).resolve().parent)
args = parser.parse_args()
args.output_dir.mkdir(parents=True, exist_ok=True)
wb = load_workbook(args.workbook, data_only=True)
gamma, recovery = 1.4, 0.89
ws = wb['Trajectory']
trajectory = np.array([(r[0], r[3], r[9], r[1]) for r in ws.iter_rows(min_row=2, values_only=True) if all(isinstance(r[i], (int,float)) for i in (0,3,9,1))], dtype=float)
trajectory = trajectory[np.argsort(trajectory[:, 0])]
assert np.isfinite(trajectory).all() and np.all(np.diff(trajectory[:,0]) > 0)
assert np.all(trajectory[:,2] > 0) and np.all(trajectory[:,1] >= 0)
ws = wb['Experiment']
sensors = []
for name, tc, vc, location in [('TS002V', 1, 2, r'$x=40.0$ mm, $\varphi=0^\circ$'), ('TS003V', 5, 6, r'$x=152.1$ mm, $\varphi=315^\circ$'), ('TS004V', 8, 9, r'$x=150.5$ mm, $\varphi=135^\circ$')]:
    data = np.array([(ws.cell(row,tc).value, ws.cell(row,vc).value) for row in range(3,ws.max_row+1) if all(isinstance(ws.cell(row,c).value,(int,float)) for c in (tc,vc))], dtype=float)
    data = data[np.argsort(data[:,0])]
    assert np.isfinite(data).all() and np.all(np.diff(data[:,0]) > 0) and np.all(data[:,1] > 0)
    sensors.append((name,data,location))
# Include every digitized knot and a fine grid for smooth evaluation of M(t)^2.
lo = max(0., trajectory[0,0], min(s[1][0,0] for s in sensors))
hi = min(trajectory[-1,0], max(s[1][-1,0] for s in sensors))
times = np.unique(np.concatenate([np.arange(lo,hi,0.1), trajectory[:,0], *[s[1][:,0] for s in sensors], [hi]]))
times = times[(times>=lo)&(times<=hi)]
mach = np.interp(times,trajectory[:,0],trajectory[:,1])
tinf = np.interp(times,trajectory[:,0],trajectory[:,2])
altitude = np.interp(times,trajectory[:,0],trajectory[:,3])
tad = tinf * (1 + recovery * (gamma-1)/2 * mach**2)
curves = {}
for name,data,_ in sensors:
    tw = np.interp(times, data[:,0],data[:,1],left=np.nan,right=np.nan)
    curves[name] = (tw,tw/tad)
with (args.output_dir/'STORT_wall_recovery_ratio.csv').open('w',newline='') as f:
    writer = csv.writer(f)
    writer.writerow(['time_s','Mach','T_inf_K','T_ad_K','altitude_m',*[v for name,_,_ in sensors for v in (name+'_T_wall_K',name+'_Tw_over_Tad')]])
    for i,t in enumerate(times):
        row = [t,mach[i],tinf[i],tad[i],altitude[i],*[v[i] for name,_,_ in sensors for v in curves[name]]]
        writer.writerow(['' if not np.isfinite(x) else format(x,'.10g') for x in row])
plt.rcParams.update({'font.size':11,'axes.spines.top':False,'axes.spines.right':False,'savefig.facecolor':'white'})
fig,ax = plt.subplots(figsize=(13,8))
for (name,data,location),color in zip(sensors,['#0072B2','#D55E00','#009E73']):
    ax.plot(times,curves[name][1],label=f'{name}  |  {location}',lw=2.3,color=color)
ax.axhline(1,color='#58606b',ls='--',lw=1.1,zorder=0)
ax.text(3,1.025,r'$T_w=T_\mathrm{ad}$',color='#58606b',fontsize=10)
ax.set(xlim=(0,215),ylim=(0,1.65),xlabel='Flight time [s]',ylabel=r'$T_w/T_\mathrm{ad}$')
ax.set_xticks(np.arange(0,201,20))
# Plot the workbook trajectory samples directly, with separate right-hand
# scales for altitude and Mach. All three axes share flight time.
trajectory_lines = []
for offset, values, label, color, style, upper, ticks in [
    (0, trajectory[:,3] / 1000., 'Altitude [km]', '#8B4C9E', '--', 45, np.arange(0,46,5)),
    (64, trajectory[:,1], 'Flight Mach number', '#454545', '-.', 10, np.arange(0,11,1)),
]:
    secondary = ax.twinx()
    secondary.set_ylim(0, upper)
    secondary.set_yticks(ticks)
    secondary.spines['right'].set_visible(True)
    secondary.spines['right'].set_position(('outward', offset))
    secondary.spines['right'].set_color(color)
    for side in ('left','top','bottom'):
        secondary.spines[side].set_visible(False)
    secondary.set_ylabel(label, color=color, labelpad=10)
    secondary.tick_params(axis='y', colors=color, labelsize=10)
    line, = secondary.plot(trajectory[:,0], values, color=color,
                           linestyle=style, linewidth=1.8, label=label)
    trajectory_lines.append(line)
    secondary.set_zorder(0)
    secondary.patch.set_visible(False)
ax.set_zorder(1)
ax.patch.set_visible(False)
ax.grid(True,alpha=.18)
fig.suptitle('STORT flight — wall-to-adiabatic recovery temperature',weight='bold',fontsize=15)
fig.text(.08,.925,r'$T_\mathrm{ad}=T_\infty[1+r(\gamma-1)M_\infty^2/2]$   |   $\gamma=1.4$, $r=0.89$',fontsize=12)
handles, labels = ax.get_legend_handles_labels()
fig.legend(handles + trajectory_lines, labels + [line.get_label() for line in trajectory_lines],
           loc='upper left', bbox_to_anchor=(.08,.90), ncol=2, frameon=False, fontsize=10)
fig.subplots_adjust(left=.08,right=.81,bottom=.14,top=.76)
fig.text(.08,.065,'Solid: temperature ratios (left axis). Dashed: altitude; dash-dot: Mach (right axes).',fontsize=9,color='#58606b')
fig.text(.08,.03,'Source: STORT_experiment_data.xlsx • Temperature ratios: 1–214 s; trajectory: 1–215 s • No extrapolation',fontsize=9,color='#58606b')
for ext in ['png','pdf','svg']:
    fig.savefig(args.output_dir/f'STORT_wall_recovery_ratio.{ext}',dpi=200)
notes = [
'# STORT wall-to-adiabatic recovery temperature', '',
f'Source: `{args.workbook}`.',
'Wall temperatures: sheet `Experiment`, columns A/B (TS002V), E/F (TS003V), H/I (TS004V), temperatures in kelvin.',
'Trajectory: sheet `Trajectory`, columns A (`time_s`), B (`altitude_m`), D (`Mach`), J (`T_inf_K`).', '',
'Assumptions: calorically perfect air, gamma = 1.4; recovery factor r = 0.89.',
'T_ad(t) = T_inf(t) * [1 + 0.178 * M_inf(t)^2].',
'The same freestream-based recovery-temperature estimate is used for all three sensors; local surface/edge-flow corrections are not included.',
'The workbook time coordinates are used directly, without shifting either dataset.',
'The left y-axis shows the three temperature ratios. Two right y-axes show altitude [km] (purple dashed) and flight Mach number (gray dash-dot) over flight time. Trajectory curves connect all 215 original Excel samples directly; they extend to 215 s.',
'Mach, freestream temperature and each wall temperature are linearly interpolated separately, then the formula and ratio are evaluated.',
'The evaluation grid includes all input time knots and a 0.1 s grid; this does not add measurement resolution.',
'Trajectory spans 1–215 s; sensors end near 214 s. Each curve is restricted to its sensor/trajectory overlap; no extrapolation.',
'Ratios greater than one mean the measured wall is hotter than the estimated instantaneous recovery temperature.', '',
'| Sensor | Last usable time [s] | Minimum ratio (time [s]) | Final ratio |',
'|---|---:|---:|---:|']
for name,data,_ in sensors:
    ratio=curves[name][1]; valid=np.flatnonzero(np.isfinite(ratio)); minimum=valid[np.argmin(ratio[valid])]; last=valid[-1]
    notes.append(f'| {name} | {times[last]:.3f} | {ratio[minimum]:.4f} ({times[minimum]:.3f}) | {ratio[last]:.4f} |')
    print(name, 'minimum',ratio[minimum], 'at', times[minimum], 'end',times[last],ratio[last], 'maximum',np.nanmax(ratio))
(args.output_dir/'README.md').write_text('\n'.join(notes)+'\n')
print('Outputs:',args.output_dir)
